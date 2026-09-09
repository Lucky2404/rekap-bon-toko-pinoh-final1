from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO


PT_COLORS = {
    "MAL": "BAE6FD",
    "SJM": "FDE68A",
    "LOGPOND": "BBF7D0",
    "TAYAN 01": "E9D5FF",
    "TAYAN 04": "FECDD3",
    "MSB": "FEE2E2",
}
DEFAULT_COLOR = "E2E8F0"

thin = Side(border_style="thin", color="94A3B8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADERS = ["No", "Toko", "No Nota", "Tanggal", "Barang", "Keterangan", "No.PP",
           "Bank", "A.n", "Rekening", "Total", "Status"]
NCOL = len(HEADERS)
TOTAL_COL = 11
LAST_LETTER = get_column_letter(NCOL)


def rupiah_format():
    return '_-"Rp" #,##0_-;-"Rp" #,##0_-;_-"Rp" "-"??_-;_-@_-'


def build_excel(periode_nama, items_by_pt):
    """items_by_pt: {pt_name: [ {toko, no_nota, tanggal, barang, keterangan, no_pp, total, bank, atas_nama, rekening, status} ]}"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Bon"

    ws["A1"] = "REKAP BON TOKO PINOH"
    ws["A1"].font = Font(size=16, bold=True, color="0F172A")
    ws.merge_cells(f"A1:{LAST_LETTER}1")
    ws["A2"] = f"Periode: {periode_nama}"
    ws["A2"].font = Font(size=11, italic=True, color="475569")
    ws.merge_cells(f"A2:{LAST_LETTER}2")

    header_row = 4
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F172A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    current_row = header_row + 1
    grand_total = 0.0

    for pt_name, items in items_by_pt.items():
        pt_color = PT_COLORS.get(pt_name.upper(), PT_COLORS.get(pt_name, DEFAULT_COLOR))
        ws.cell(row=current_row, column=1, value=f"PT {pt_name}")
        ws.cell(row=current_row, column=1).font = Font(bold=True, size=12, color="0F172A")
        ws.cell(row=current_row, column=1).fill = PatternFill("solid", fgColor=pt_color)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=NCOL)
        current_row += 1

        pt_total = 0.0
        by_toko = {}
        for it in items:
            by_toko.setdefault(it["toko"], []).append(it)

        row_no = 1
        for toko, toko_items in by_toko.items():
            toko_total = 0.0
            for it in toko_items:
                ws.cell(row=current_row, column=1, value=row_no)
                ws.cell(row=current_row, column=2, value=it["toko"])
                ws.cell(row=current_row, column=3, value=it["no_nota"])
                ws.cell(row=current_row, column=4, value=it["tanggal"])
                ws.cell(row=current_row, column=5, value=it.get("barang", ""))
                ws.cell(row=current_row, column=6, value=it["keterangan"])
                ws.cell(row=current_row, column=7, value=it["no_pp"])
                ws.cell(row=current_row, column=8, value=it["bank"])
                ws.cell(row=current_row, column=9, value=it["atas_nama"])
                ws.cell(row=current_row, column=10, value=str(it["rekening"]))
                tc = ws.cell(row=current_row, column=TOTAL_COL, value=it["total"])
                tc.number_format = rupiah_format()
                tc.alignment = Alignment(horizontal="right")
                ws.cell(row=current_row, column=12, value=it["status"])
                for col in range(1, NCOL + 1):
                    ws.cell(row=current_row, column=col).border = BORDER
                toko_total += float(it["total"] or 0)
                row_no += 1
                current_row += 1

            ws.cell(row=current_row, column=2, value=f"Subtotal {toko}").font = Font(bold=True, italic=True)
            tc = ws.cell(row=current_row, column=TOTAL_COL, value=toko_total)
            tc.font = Font(bold=True)
            tc.number_format = rupiah_format()
            tc.alignment = Alignment(horizontal="right")
            for col in range(1, NCOL + 1):
                ws.cell(row=current_row, column=col).fill = PatternFill("solid", fgColor="F1F5F9")
                ws.cell(row=current_row, column=col).border = BORDER
            current_row += 1
            pt_total += toko_total

        ws.cell(row=current_row, column=2, value=f"TOTAL PT {pt_name}").font = Font(bold=True, color="0F172A")
        tc = ws.cell(row=current_row, column=TOTAL_COL, value=pt_total)
        tc.font = Font(bold=True, color="0F172A")
        tc.number_format = rupiah_format()
        tc.alignment = Alignment(horizontal="right")
        for col in range(1, NCOL + 1):
            ws.cell(row=current_row, column=col).fill = PatternFill("solid", fgColor=pt_color)
            ws.cell(row=current_row, column=col).border = BORDER
        current_row += 2
        grand_total += pt_total

    ws.cell(row=current_row, column=2, value="GRAND TOTAL").font = Font(bold=True, size=13, color="FFFFFF")
    tc = ws.cell(row=current_row, column=TOTAL_COL, value=grand_total)
    tc.font = Font(bold=True, size=13, color="FFFFFF")
    tc.number_format = rupiah_format()
    tc.alignment = Alignment(horizontal="right")
    for col in range(1, NCOL + 1):
        ws.cell(row=current_row, column=col).fill = PatternFill("solid", fgColor="0F172A")

    widths = [5, 22, 14, 12, 26, 30, 12, 14, 22, 20, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:{LAST_LETTER}{header_row}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
